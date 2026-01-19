"""Template-based Excel export service for sub-scheme 20290182.

Architecture mirrors 20290046 with hardcoded row mappings for merged-cell templates.
"""
import io
import os
from typing import Optional
from copy import copy

from fastapi import HTTPException
from starlette.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import load_workbook, Workbook

from ..config import SHEET_NAMES, SCHEME_DISTRICTS
from src.schemes.common.excel_export import ExcelExportService
from .populators.budget_post_details import populate_budget_post_details
from .populators.post_status import populate_post_status
from .populators.post_expenses import populate_post_expenses
from .populators.unit_expenditure import populate_unit_expenditure

DISTRICT_PROCESSORS = {
    "Mumbai City": "mumbai_city",
    "Mumbai Suburban": "mumbai_suburban",
    "Thane": "thane",
    "Palghar": "palghar",
    "Raigad": "raigad",
    "Ratnagiri": "ratnagiri",
    "Sindhudurg": "sindhudurg",
    "DCO Staff": "dco_staff",
}


def _get_template_path(sub_scheme_code: Optional[str] = None) -> str:
    if sub_scheme_code:
        parent_scheme = sub_scheme_code[:4]
        template_dir = f"excel_templates/s{parent_scheme}/subs/s{sub_scheme_code}"
        for ext in (".xlsx", ".xls"):
            for pattern in (
                f"Budget {sub_scheme_code} for 2026-27{ext}",
                f"original_template{ext}",
            ):
                path = f"{template_dir}/{pattern}"
                if os.path.exists(path):
                    return path
    return "excel_templates/s2029/subs/s20290182/Budget 20290182 for 2026-27.xlsx"


def _get_processor_module(district: str):
    processor_name = DISTRICT_PROCESSORS.get(district)
    if not processor_name:
        return None
    try:
        return __import__(
            f"src.schemes.s2029.subs.s20290182.excel_export.processors.{processor_name}",
            fromlist=[processor_name],
        )
    except ImportError:
        from .processors import mumbai_city

        return mumbai_city


def _copy_sheet_to_new_workbook(source_sheet, sheet_name: str) -> Workbook:
    new_wb = Workbook()
    new_wb.remove(new_wb.active)
    target_sheet = new_wb.create_sheet(sheet_name)
    for row in source_sheet.iter_rows():
        for cell in row:
            if not hasattr(cell, "coordinate"):
                continue
            target_cell = target_sheet[cell.coordinate]
            target_cell.value = cell.value
            if hasattr(cell, "has_style") and cell.has_style:
                target_cell.font = copy(cell.font)
                target_cell.border = copy(cell.border)
                target_cell.fill = copy(cell.fill)
                target_cell.number_format = cell.number_format
                target_cell.protection = copy(cell.protection)
                target_cell.alignment = copy(cell.alignment)
    for merged_range in source_sheet.merged_cells.ranges:
        target_sheet.merge_cells(str(merged_range))
    for col_letter, col_dim in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[col_letter].width = col_dim.width
    for row_num, row_dim in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[row_num].height = row_dim.height
    return new_wb


def _generate_workbook(
    db: Session,
    only_sheet: Optional[str],
    user_district: Optional[str],
    sub_scheme_code: Optional[str],
    fiscal_year: Optional[str],
) -> io.BytesIO:
    template_path = _get_template_path(sub_scheme_code)
    wb = load_workbook(template_path, data_only=False)

    if only_sheet in (None, "budget_post_details"):
        populate_budget_post_details(wb, db, sub_scheme_code, fiscal_year)
    if only_sheet in (None, "post_status"):
        populate_post_status(wb, db, sub_scheme_code, fiscal_year)
    if only_sheet in (None, "post_expenses"):
        populate_post_expenses(wb, db, sub_scheme_code, fiscal_year)
    if only_sheet in (None, "unit_expenditure"):
        populate_unit_expenditure(wb, db, sub_scheme_code, fiscal_year)

    if user_district is not None and user_district in SCHEME_DISTRICTS:
        wb = _apply_district_filter(wb, user_district, only_sheet)
    elif only_sheet is not None:
        sheet_name = SHEET_NAMES.get(only_sheet)
        if sheet_name and sheet_name in wb.sheetnames:
            wb = _copy_sheet_to_new_workbook(wb[sheet_name], sheet_name)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _apply_district_filter(
    wb: Workbook, district: str, only_sheet: Optional[str]
) -> Workbook:
    processor = _get_processor_module(district)
    if not processor:
        from .processors import mumbai_city as processor

    apply_filtering = processor.apply_district_filtering
    sheets_to_exclude = getattr(processor, "SHEETS_TO_EXCLUDE", [])

    if only_sheet is not None:
        sheet_name = SHEET_NAMES.get(only_sheet)
        if sheet_name and sheet_name in wb.sheetnames:
            source_sheet = wb[sheet_name]
            apply_filtering(source_sheet, only_sheet)
            return _copy_sheet_to_new_workbook(source_sheet, sheet_name)
        return wb

    for name in [n for n in wb.sheetnames if n in sheets_to_exclude]:
        wb.remove(wb[name])
    for sheet_key, sheet_name in SHEET_NAMES.items():
        if sheet_name in wb.sheetnames:
            apply_filtering(wb[sheet_name], sheet_key)
    return wb


async def export_original_workbook_async(
    db: Session,
    only_sheet: Optional[str] = None,
    user_district: Optional[str] = None,
    sub_scheme_code: Optional[str] = None,
    fiscal_year: Optional[str] = None,
) -> StreamingResponse:
    base_filename = (
        "original_format" if only_sheet is None else f"{only_sheet}_original_format"
    )

    def generate():
        try:
            return _generate_workbook(
                db, only_sheet, user_district, sub_scheme_code, fiscal_year
            )
        except Exception as exc:  # pragma: no cover - defensive
            raise HTTPException(
                status_code=500, detail=f"Failed to generate Excel: {exc}"
            )

    return await ExcelExportService.export_with_throttle(
        export_fn=generate, filename=base_filename, fiscal_year=fiscal_year
    )


def export_original_workbook(
    db: Session,
    only_sheet: Optional[str] = None,
    user_district: Optional[str] = None,
    sub_scheme_code: Optional[str] = None,
    fiscal_year: Optional[str] = None,
) -> StreamingResponse:
    try:
        output = _generate_workbook(
            db, only_sheet, user_district, sub_scheme_code, fiscal_year
        )
    except Exception as exc:  # pragma: no cover - defensive
        raise HTTPException(
            status_code=500, detail=f"Failed to generate Excel: {exc}"
        )
    base_filename = (
        "original_format" if only_sheet is None else f"{only_sheet}_original_format"
    )
    return ExcelExportService.create_response(
        content=output, base_filename=base_filename, fiscal_year=fiscal_year
    )

