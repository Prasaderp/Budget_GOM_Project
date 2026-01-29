"""Unified Excel export for s2045 district expenditure sub-schemes.

All 3 sub-schemes (20450182, 20450251, 20450262) are in the SAME Excel file (Sheet 1).
This service exports data for ALL sub-schemes together regardless of which sub-scheme
initiated the export.

Excel Layout:
- Rows 7-13: 20450182 (7 districts)
- Rows 25-29: 20450251 (5 districts, no Mumbai)
- Rows 41-47: 20450262 (7 districts)
"""
import io
import os
from typing import Optional, Dict, Any

from fastapi import HTTPException
from starlette.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from src.schemes.common.excel_export import ExcelExportService
from .excel_populator import (
    populate_district_expenditure,
    ROW_MAP_20450182,
    ROW_MAP_20450251,
    ROW_MAP_20450262,
)

# Single template location for ALL 3 sub-schemes
TEMPLATE_DIR = "excel_templates/s2045/subs/s2045_district_expenditure"
SHEET_NAME = "Sheet1"

# Sub-scheme codes
SUBSCHEME_20450182 = "20450182"
SUBSCHEME_20450251 = "20450251"
SUBSCHEME_20450262 = "20450262"


def _get_template_path() -> Optional[str]:
    """Find the Excel template file in the shared template directory."""
    if not os.path.exists(TEMPLATE_DIR):
        return None
    
    for filename in os.listdir(TEMPLATE_DIR):
        if filename.endswith(('.xlsx', '.xls')):
            return os.path.join(TEMPLATE_DIR, filename)
    
    return None


def _get_model_classes():
    """
    Lazy import of model classes to avoid circular imports.
    Returns tuple of (Model20450182, Model20450251, Model20450262).
    """
    from src.schemes.s2045.subs.s20450182.models import DistrictExpenditure20450182
    from src.schemes.s2045.subs.s20450251.models import DistrictExpenditure20450251
    from src.schemes.s2045.subs.s20450262.models import DistrictExpenditure20450262
    
    return (
        DistrictExpenditure20450182,
        DistrictExpenditure20450251,
        DistrictExpenditure20450262,
    )


def _generate_unified_workbook(
    db: Session,
    fiscal_year: Optional[str],
) -> io.BytesIO:
    """
    Load template and populate with data from ALL 3 sub-schemes.
    
    This ensures consistent export regardless of which sub-scheme initiated
    the download request.
    """
    template_path = _get_template_path()
    if not template_path:
        raise HTTPException(
            status_code=404,
            detail=f"Excel template not found in {TEMPLATE_DIR}. "
                   "Please ensure the template file exists.",
        )
    
    try:
        wb = load_workbook(template_path, data_only=False)
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to load Excel template: {e}",
        )
    
    # Get worksheet
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    
    # Get model classes
    Model20450182, Model20450251, Model20450262 = _get_model_classes()
    
    # Query and populate data for ALL 3 sub-schemes
    
    # 1. Populate 20450182 (7 Konkan districts, rows 7-13)
    query_182 = db.query(Model20450182).filter(
        Model20450182.sub_scheme_code == SUBSCHEME_20450182
    )
    if fiscal_year:
        query_182 = query_182.filter(Model20450182.fiscal_year == fiscal_year)
    records_182 = query_182.all()
    populate_district_expenditure(ws, records_182, ROW_MAP_20450182)
    
    # 2. Populate 20450251 (5 districts - no Mumbai, rows 25-29)
    query_251 = db.query(Model20450251).filter(
        Model20450251.sub_scheme_code == SUBSCHEME_20450251
    )
    if fiscal_year:
        query_251 = query_251.filter(Model20450251.fiscal_year == fiscal_year)
    records_251 = query_251.all()
    populate_district_expenditure(ws, records_251, ROW_MAP_20450251)
    
    # 3. Populate 20450262 (7 Konkan districts, rows 41-47)
    query_262 = db.query(Model20450262).filter(
        Model20450262.sub_scheme_code == SUBSCHEME_20450262
    )
    if fiscal_year:
        query_262 = query_262.filter(Model20450262.fiscal_year == fiscal_year)
    records_262 = query_262.all()
    populate_district_expenditure(ws, records_262, ROW_MAP_20450262)
    
    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def export_unified_workbook_async(
    db: Session,
    fiscal_year: Optional[str] = None,
) -> StreamingResponse:
    """
    Async export with throttling support.
    
    This is the MAIN export function to be used by all 3 sub-schemes.
    Exports data for ALL sub-schemes (20450182, 20450251, 20450262) together.
    """
    def generate():
        try:
            return _generate_unified_workbook(db, fiscal_year)
        except Exception as exc:
            if isinstance(exc, HTTPException):
                raise exc
            raise HTTPException(
                status_code=500,
                detail=f"Failed to generate Excel: {exc}"
            )
    
    return await ExcelExportService.export_with_throttle(
        export_fn=generate,
        filename="2045_district_expenditure_all",
        fiscal_year=fiscal_year,
    )


def export_unified_workbook(
    db: Session,
    fiscal_year: Optional[str] = None,
) -> StreamingResponse:
    """Synchronous version for direct calls."""
    try:
        output = _generate_unified_workbook(db, fiscal_year)
    except Exception as exc:
        if isinstance(exc, HTTPException):
            raise exc
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate Excel: {exc}"
        )
    
    return ExcelExportService.create_response(
        content=output,
        base_filename="2045_district_expenditure_all",
        fiscal_year=fiscal_year,
    )
