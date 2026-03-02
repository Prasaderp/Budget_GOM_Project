"""Populator for scheme 2245 data into combined workbook.

Handles both Section 1 (district-wise) and Section 3 (DC/ZP splits) data.
"""
from typing import Optional
from sqlalchemy.orm import Session
from openpyxl.workbook import Workbook

from src.schemes.s2245.subs.s2245.models import DistrictExpenditure2245, SUB_SCHEME_CODE
from src.schemes.s2245.subs.s2245.config import TABLE_SECTIONS, EXTRA_DISTRICT, get_section3_table_sections

SHEET_SECTION1 = "155 217 271 315  2461"
SHEET_SECTION3 = "22450093 22452185"

SECTION1_COL_MAP = {
    "exp_prev3": "C",
    "exp_prev2": "D",
    "exp_prev1": "E",
    "budget_estimate_curr": "F",
    "revised_estimate_curr": "G",
    "budget_estimate_next": "H",
}

SECTION3_COL_MAP = {
    "exp_prev3": "D",
    "exp_prev2": "E",
    "exp_prev1": "F",
    "budget_estimate_curr": "G",
    "revised_estimate_curr": "H",
    "budget_estimate_next": "I",
}

TABLE_CODES = [s["code"] for s in TABLE_SECTIONS if not s.get("is_section3")]

SECTION1_ROW_MAP = {
    "22450155": {"Mumbai City": 8, "Mumbai Suburban": 9, "Thane": 10, "Palghar": 11, "Raigad": 12, "Ratnagiri": 13, "Sindhudurg": 14},
    "22450182": {"Mumbai City": 17, "Mumbai Suburban": 18, "Thane": 19, "Palghar": 20, "Raigad": 21, "Ratnagiri": 22, "Sindhudurg": 23},
    "22450191": {"Mumbai City": 26, "Mumbai Suburban": 27, "Thane": 28, "Palghar": 29, "Raigad": 30, "Ratnagiri": 31, "Sindhudurg": 32},
    "22450217": {"Mumbai City": 35, "Mumbai Suburban": 36, "Thane": 37, "Palghar": 38, "Raigad": 39, "Ratnagiri": 40, "Sindhudurg": 41},
    "22450244": {"Mumbai City": 44, "Mumbai Suburban": 45, "Thane": 46, "Palghar": 47, "Raigad": 48, "Ratnagiri": 49, "Sindhudurg": 50},
    "22450271": {"Mumbai City": 53, "Mumbai Suburban": 54, "Thane": 55, "Palghar": 56, "Raigad": 57, "Ratnagiri": 58, "Sindhudurg": 59},
    "22450291": {"Mumbai City": 62, "Mumbai Suburban": 63, "Thane": 64, "Palghar": 65, "Raigad": 66, "Ratnagiri": 67, "Sindhudurg": 68},
    "22450315": {"Mumbai City": 71, "Mumbai Suburban": 72, "Thane": 73, "Palghar": 74, "Raigad": 75, "Ratnagiri": 76, "Sindhudurg": 77},
    "22450324": {"Mumbai City": 80, "Mumbai Suburban": 81, "Thane": 82, "Palghar": 83, "Raigad": 84, "Ratnagiri": 85, "Sindhudurg": 86},
    "22450333": {"Mumbai City": 89, "Mumbai Suburban": 90, "Thane": 91, "Palghar": 92, "Raigad": 93, "Ratnagiri": 94, "Sindhudurg": 95},
    "22450988": {"Mumbai City": 98, "Mumbai Suburban": 99, "Thane": 100, "Palghar": 101, "Raigad": 102, "Ratnagiri": 103, "Sindhudurg": 104},
    "22452194": {"Mumbai City": 107, "Mumbai Suburban": 108, "Thane": 109, "Palghar": 110, "Raigad": 111, "Ratnagiri": 112, "Sindhudurg": 113},
    "22452247": {"Mumbai City": 116, "Mumbai Suburban": 117, "Thane": 118, "Palghar": 119, "Raigad": 120, "Ratnagiri": 121, "Sindhudurg": 122},
    "22452309": {"Mumbai City": 125, "Mumbai Suburban": 126, "Thane": 127, "Palghar": 128, "Raigad": 129, "Ratnagiri": 130, "Sindhudurg": 131},
    "22452327": {"Mumbai City": 134, "Mumbai Suburban": 135, "Thane": 136, "Palghar": 137, "Raigad": 138, "Ratnagiri": 139, "Sindhudurg": 140},
    "22452363": {"Mumbai City": 143, "Mumbai Suburban": 144, "Thane": 145, "Palghar": 146, "Raigad": 147, "Ratnagiri": 148, "Sindhudurg": 149},
    "22452372": {"Mumbai City": 152, "Mumbai Suburban": 153, "Thane": 154, "Palghar": 155, "Raigad": 156, "Ratnagiri": 157, "Sindhudurg": 158},
    "22452381": {"Mumbai City": 161, "Mumbai Suburban": 162, "Thane": 163, "Palghar": 164, "Raigad": 165, "Ratnagiri": 166, "Sindhudurg": 167},
    "22452407": {EXTRA_DISTRICT: 170, "Mumbai City": 171, "Mumbai Suburban": 172, "Thane": 173, "Palghar": 174, "Raigad": 175, "Ratnagiri": 176, "Sindhudurg": 177},
    "22452434": {"Mumbai City": 180, "Mumbai Suburban": 181, "Thane": 182, "Palghar": 183, "Raigad": 184, "Ratnagiri": 185, "Sindhudurg": 186},
    "22452452": {"Mumbai City": 189, "Mumbai Suburban": 190, "Thane": 191, "Palghar": 192, "Raigad": 193, "Ratnagiri": 194, "Sindhudurg": 195},
    "22452461": {"Mumbai City": 198, "Mumbai Suburban": 199, "Thane": 200, "Palghar": 201, "Raigad": 202, "Ratnagiri": 203, "Sindhudurg": 204},
    "22452472": {"Mumbai City": 207, "Mumbai Suburban": 208, "Thane": 209, "Palghar": 210, "Raigad": 211, "Ratnagiri": 212, "Sindhudurg": 213},
    "22452499": {"Mumbai City": 216, "Mumbai Suburban": 217, "Thane": 218, "Palghar": 219, "Raigad": 220, "Ratnagiri": 221, "Sindhudurg": 222},
    "22452603": {"Mumbai City": 225, "Mumbai Suburban": 226, "Thane": 227, "Palghar": 228, "Raigad": 229, "Ratnagiri": 230, "Sindhudurg": 231},
    "22454141": {"Mumbai City": 235, "Mumbai Suburban": 236, "Thane": 237, "Palghar": 238, "Raigad": 239, "Ratnagiri": 240, "Sindhudurg": 241, EXTRA_DISTRICT: 242},
    "22454188": {"Mumbai City": 245, "Mumbai Suburban": 246, "Thane": 247, "Palghar": 248, "Raigad": 249, "Ratnagiri": 250, "Sindhudurg": 251, EXTRA_DISTRICT: 252},
    "22451761_10": {EXTRA_DISTRICT: 255, "Mumbai City": 256, "Mumbai Suburban": 257, "Thane": 258, "Palghar": 259, "Raigad": 260, "Ratnagiri": 261, "Sindhudurg": 262},
    "22451761_11": {EXTRA_DISTRICT: 265, "Mumbai City": 266, "Mumbai Suburban": 267, "Thane": 268, "Palghar": 269, "Raigad": 270, "Ratnagiri": 271, "Sindhudurg": 272},
    "22451761_21": {EXTRA_DISTRICT: 275, "Mumbai City": 276, "Mumbai Suburban": 277, "Thane": 278, "Palghar": 279, "Raigad": 280, "Ratnagiri": 281, "Sindhudurg": 282},
    "22451761_27": {EXTRA_DISTRICT: 285, "Mumbai City": 286, "Mumbai Suburban": 287, "Thane": 288, "Palghar": 289, "Raigad": 290, "Ratnagiri": 291, "Sindhudurg": 292},
    "22451761_31": {EXTRA_DISTRICT: 295, "Mumbai City": 296, "Mumbai Suburban": 297, "Thane": 298, "Palghar": 299, "Raigad": 300, "Ratnagiri": 301, "Sindhudurg": 302},
    "22451761_52": {EXTRA_DISTRICT: 305, "Mumbai City": 306, "Mumbai Suburban": 307, "Thane": 308, "Palghar": 309, "Raigad": 310, "Ratnagiri": 311, "Sindhudurg": 312},
}

SECTION3_ROW_MAP = {
    "22450093": {
        "Thane": {"DC": 8, "ZP": 9},
        "Palghar": {"DC": 11, "ZP": 12},
        "Raigad": {"DC": 15, "ZP": 16},
        "Ratnagiri": {"DC": 18, "ZP": 19},
        "Sindhudurg": {"DC": 21, "ZP": 22},
    },
    "22452185": {
        "Thane": {"DC": 25, "ZP": 26},
        "Palghar": {"DC": 28, "ZP": 29},
        "Raigad": {"DC": 32, "ZP": 33},
        "Ratnagiri": {"DC": 35, "ZP": 36},
        "Sindhudurg": {"DC": 38, "ZP": 39},
    },
}


def _populate_section1(wb: Workbook, db: Session, fiscal_year: Optional[str]) -> None:
    """Populate section 1 sheet with district expenditure data."""
    if SHEET_SECTION1 not in wb.sheetnames:
        return
    ws = wb[SHEET_SECTION1]

    query = db.query(DistrictExpenditure2245).filter(
        DistrictExpenditure2245.sub_scheme_code == SUB_SCHEME_CODE,
        DistrictExpenditure2245.table_section_code.in_(TABLE_CODES),
    )
    if fiscal_year:
        query = query.filter(DistrictExpenditure2245.fiscal_year == fiscal_year)

    records = query.all()
    data_map = {(r.table_section_code, r.district): r for r in records}

    for table_code, district_rows in SECTION1_ROW_MAP.items():
        for district, row_num in district_rows.items():
            if row_num == 0:
                continue
            record = data_map.get((table_code, district))
            if not record:
                continue
            for field, col in SECTION1_COL_MAP.items():
                ws[f"{col}{row_num}"] = getattr(record, field, 0) or 0


def _populate_section3(wb: Workbook, db: Session, fiscal_year: Optional[str]) -> None:
    """Populate section 3 sheet with DC/ZP split data."""
    if SHEET_SECTION3 not in wb.sheetnames:
        return
    ws = wb[SHEET_SECTION3]

    section3_codes = [s["code"] for s in get_section3_table_sections()]
    if not section3_codes:
        return

    query = db.query(DistrictExpenditure2245).filter(
        DistrictExpenditure2245.sub_scheme_code == SUB_SCHEME_CODE,
        DistrictExpenditure2245.table_section_code.in_(section3_codes),
    )
    if fiscal_year:
        query = query.filter(DistrictExpenditure2245.fiscal_year == fiscal_year)

    records = query.all()
    data_map = {(r.table_section_code, r.district): r for r in records}

    for table_code, district_data in SECTION3_ROW_MAP.items():
        for district, row_types in district_data.items():
            for row_type, row_num in row_types.items():
                key = (table_code, f"{district}|{row_type}")
                record = data_map.get(key)
                if not record:
                    continue
                for field, col in SECTION3_COL_MAP.items():
                    ws[f"{col}{row_num}"] = getattr(record, field, 0) or 0


def populate_s2245_data(wb: Workbook, db: Session, fiscal_year: Optional[str]) -> None:
    """Populate all 2245 scheme sheets in workbook."""
    _populate_section1(wb, db, fiscal_year)
    _populate_section3(wb, db, fiscal_year)
