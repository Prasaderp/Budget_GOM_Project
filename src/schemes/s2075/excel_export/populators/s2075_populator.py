"""Unified populator for scheme 2075 Excel export.

Handles both sub-schemas across 3 sheets in a single Excel file:
- Sheet "2075 2019-20": Master summary with both sub-schemes
- Sheet "249": 20750249 detail (Single Entry - DCO only)
- Sheet "2075 294": 20750294 detail (4 districts: Thane, Palghar, Raigad, Sindhudurg)
"""
from typing import Optional, Dict, Any
from sqlalchemy.orm import Session
from sqlalchemy import func
from openpyxl.workbook import Workbook

from src.schemes.s2075.models import SubHeadExpenditure2075, DistrictExpenditure2075
from src.schemes.s2075.config import DISTRICTS, SHEET_NAMES

# Column mapping for all sheets (same structure)
COL_MAP = {
    "expenditure_prev3": "C",
    "expenditure_prev2": "D",
    "expenditure_prev1": "E",
    "budget_estimate_curr": "F",
    "revised_estimate_curr": "G",
    "budget_estimate_next": "H",
}

# Master sheet row mapping
MASTER_ROW_249 = 9
MASTER_ROW_294 = 10

# Detail sheet row mappings (based on actual Excel structure)
ROW_249 = 8

# District row mapping for sheet "2075 294"
DISTRICT_ROW_MAP_294 = {
    "Thane": 9,
    "Palghar": 10,
    "Raigad": 11,
    "Sindhudurg": 12,
}

FIELDS = list(COL_MAP.keys())
INTEGER_FORMAT = '0'


# ============================================================================
# UTILITY FUNCTIONS
# ============================================================================

def _ensure_integer(value: Any) -> int:
    """Convert any numeric value to integer, handling None and zero cases."""
    if value is None:
        return 0
    return int(value)


def _write_integer_cell(ws: Any, cell_address: str, value: Any) -> None:
    """Write an integer value to an Excel cell with proper formatting."""
    cell = ws[cell_address]
    cell.value = _ensure_integer(value)
    cell.number_format = INTEGER_FORMAT


# ============================================================================
# DATA FETCHING
# ============================================================================

def _fetch_249_data(db: Session, fiscal_year: Optional[str]) -> Optional[Any]:
    """Fetch single row data for 20750249."""
    query = db.query(SubHeadExpenditure2075).filter(
        SubHeadExpenditure2075.sub_scheme_code == "20750249",
    )
    if fiscal_year:
        query = query.filter(SubHeadExpenditure2075.fiscal_year == fiscal_year)
    return query.first()


def _fetch_294_data(db: Session, fiscal_year: Optional[str]) -> Dict[str, Any]:
    """Fetch district data for 20750294 as district->record map."""
    query = db.query(DistrictExpenditure2075).filter(
        DistrictExpenditure2075.sub_scheme_code == "20750294",
    )
    if fiscal_year:
        query = query.filter(DistrictExpenditure2075.fiscal_year == fiscal_year)
    records = query.all()
    return {r.district: r for r in records}


def _fetch_294_aggregated(db: Session, fiscal_year: Optional[str]) -> Dict[str, int]:
    """Fetch aggregated totals for 20750294 (sum of all districts)."""
    query = db.query(
        func.coalesce(func.sum(DistrictExpenditure2075.expenditure_prev3), 0),
        func.coalesce(func.sum(DistrictExpenditure2075.expenditure_prev2), 0),
        func.coalesce(func.sum(DistrictExpenditure2075.expenditure_prev1), 0),
        func.coalesce(func.sum(DistrictExpenditure2075.budget_estimate_curr), 0),
        func.coalesce(func.sum(DistrictExpenditure2075.revised_estimate_curr), 0),
        func.coalesce(func.sum(DistrictExpenditure2075.budget_estimate_next), 0),
    ).filter(DistrictExpenditure2075.sub_scheme_code == "20750294")
    
    if fiscal_year:
        query = query.filter(DistrictExpenditure2075.fiscal_year == fiscal_year)
    
    result = query.first()
    if not result:
        return {f: 0 for f in FIELDS}
    
    return {
        "expenditure_prev3": _ensure_integer(result[0]),
        "expenditure_prev2": _ensure_integer(result[1]),
        "expenditure_prev1": _ensure_integer(result[2]),
        "budget_estimate_curr": _ensure_integer(result[3]),
        "revised_estimate_curr": _ensure_integer(result[4]),
        "budget_estimate_next": _ensure_integer(result[5]),
    }


# ============================================================================
# ROW POPULATION
# ============================================================================

def _populate_row_from_record(ws: Any, row: int, record: Any) -> None:
    """Populate a single row with expenditure data from a model record."""
    for field, col in COL_MAP.items():
        raw_value = getattr(record, field, None)
        _write_integer_cell(ws, f"{col}{row}", raw_value)


def _populate_row_from_dict(ws: Any, row: int, data: Dict[str, int]) -> None:
    """Populate a single row with expenditure data from a dictionary."""
    for field, col in COL_MAP.items():
        raw_value = data.get(field)
        _write_integer_cell(ws, f"{col}{row}", raw_value)


# ============================================================================
# SHEET POPULATION
# ============================================================================

def _populate_master(wb: Workbook, db: Session, fiscal_year: Optional[str]) -> None:
    """Populate master summary sheet with both sub-schemes."""
    sheet_name = SHEET_NAMES["master"]
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]
    
    # Row 9: 20750249 data (single entry)
    record_249 = _fetch_249_data(db, fiscal_year)
    if record_249:
        _populate_row_from_record(ws, MASTER_ROW_249, record_249)
    
    # Row 10: 20750294 aggregated data (sum of all districts)
    data_294 = _fetch_294_aggregated(db, fiscal_year)
    _populate_row_from_dict(ws, MASTER_ROW_294, data_294)


def _populate_249(wb: Workbook, db: Session, fiscal_year: Optional[str]) -> None:
    """Populate 20750249 detail sheet (single row, DCO only)."""
    sheet_name = SHEET_NAMES["sub_head_249"]
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]
    record = _fetch_249_data(db, fiscal_year)
    if record:
        _populate_row_from_record(ws, ROW_249, record)


def _populate_294(wb: Workbook, db: Session, fiscal_year: Optional[str]) -> None:
    """Populate 20750294 detail sheet (4 districts)."""
    sheet_name = SHEET_NAMES["district_294"]
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]
    data_map = _fetch_294_data(db, fiscal_year)
    for district in DISTRICTS:
        row = DISTRICT_ROW_MAP_294.get(district)
        if not row:
            continue
        record = data_map.get(district)
        if record:
            _populate_row_from_record(ws, row, record)


# ============================================================================
# MAIN ENTRY POINT
# ============================================================================

def populate_s2075_data(
    wb: Workbook,
    db: Session,
    fiscal_year: Optional[str],
) -> None:
    """Populate all 2075 sheets in the workbook."""
    _populate_master(wb, db, fiscal_year)
    _populate_249(wb, db, fiscal_year)
    _populate_294(wb, db, fiscal_year)
