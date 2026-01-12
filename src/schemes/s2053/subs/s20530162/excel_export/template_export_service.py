"""Template-based Excel export service for sub-scheme 20530162.

This module handles the export of original Excel template workbooks
with data population and district filtering. Uses centralized
ExcelExportService for production-grade throttling and caching.

IMPORTANT: For production endpoints with 500+ concurrent users,
use export_original_workbook_async which includes throttling.
"""
import io
import os
from typing import Optional
from copy import copy

from fastapi import HTTPException
from starlette.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import load_workbook, Workbook

from ..config import SHEET_NAMES
from src.schemes.common.excel_export import ExcelExportService, get_no_cache_headers, build_filename
from .populators.budget_post_details import populate_budget_post_details
from .populators.post_status import populate_post_status
from .populators.post_expenses import populate_post_expenses
from .populators.unit_expenditure import populate_unit_expenditure


def _get_template_path(sub_scheme_code: Optional[str] = None) -> str:
    """Get the path to the Excel template file."""
    if sub_scheme_code:
        parent_scheme = sub_scheme_code[:4]
        # Try .xls first (old format)
        template_path_xls = f"excel_templates/s{parent_scheme}/subs/s{sub_scheme_code}/Budget {sub_scheme_code} for 2026-27.xls"
        if os.path.exists(template_path_xls):
            return template_path_xls
        # Fallback to .xlsx
        template_path_xlsx = f"excel_templates/s{parent_scheme}/subs/s{sub_scheme_code}/original_template.xlsx"
        if os.path.exists(template_path_xlsx):
            return template_path_xlsx
    # Default fallback
    return "excel_templates/s2053/subs/s20530162/Budget 20530162 for 2026-27.xlsx"


def _get_processor_module(district: str):
    """Get the district-specific processor module."""
    DISTRICT_PROCESSORS = {
        "Mumbai City": "mumbai_city",
        "Mumbai Suburban": "mumbai_suburban",
        "Thane": "thane",
        "Palghar": "palghar",
        "Raigad": "raigad",
        "Ratnagiri": "ratnagiri",
        "Sindhudurg": "sindhudurg",
        "DCO Staff": "dco_staff" 
    }
    
    processor_name = DISTRICT_PROCESSORS.get(district)
    if not processor_name:
        return None
    
    try:
        processor_module = __import__(
            f'src.schemes.s2053.subs.s20530162.excel_export.processors.{processor_name}',
            fromlist=[processor_name]
        )
        return processor_module
    except ImportError:
        from .processors import mumbai_city
        return mumbai_city


def _copy_sheet_to_new_workbook(source_sheet, sheet_name: str) -> Workbook:
    """Copy a single sheet to a new workbook with styles preserved."""
    new_wb = Workbook()
    new_wb.remove(new_wb.active)
    target_sheet = new_wb.create_sheet(sheet_name)
    
    for row in source_sheet.iter_rows():
        for cell in row:
            if hasattr(cell, 'coordinate'):
                target_cell = target_sheet[cell.coordinate]
                target_cell.value = cell.value
                if hasattr(cell, 'has_style') and cell.has_style:
                    target_cell.font = copy(cell.font)
                    target_cell.border = copy(cell.border)
                    target_cell.fill = copy(cell.fill)
                    target_cell.number_format = cell.number_format
                    target_cell.protection = copy(cell.protection)
                    target_cell.alignment = copy(cell.alignment)
    
    for merged_range in source_sheet.merged_cells.ranges:
        target_sheet.merge_cells(str(merged_range))
    
    for col_letter, col_dim in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[col_letter].width = col_dim.width
    
    for row_num, row_dim in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[row_num].height = row_dim.height
    
    return new_wb


def _generate_workbook(
    db: Session,
    only_sheet: Optional[str],
    user_district: Optional[str],
    sub_scheme_code: Optional[str],
    fiscal_year: Optional[str]
) -> io.BytesIO:
    """
    Internal function to generate populated workbook.
    
    This is the CPU-intensive work that runs in the thread pool.
    """
    template_path = _get_template_path(sub_scheme_code)
    wb = load_workbook(template_path, data_only=False)
    
    # Populate sheets with data
    if only_sheet in (None, "budget_post_details"):
        populate_budget_post_details(wb, db, sub_scheme_code, fiscal_year)
    if only_sheet in (None, "post_status"):
        populate_post_status(wb, db, sub_scheme_code, fiscal_year)
    if only_sheet in (None, "post_expenses"):
        populate_post_expenses(wb, db, sub_scheme_code, fiscal_year)
    if only_sheet in (None, "unit_expenditure"):
        populate_unit_expenditure(wb, db, sub_scheme_code, fiscal_year)

    # Apply district-specific filtering
    if user_district is not None:
        processor_module = _get_processor_module(user_district)
        if processor_module:
            apply_district_filtering = processor_module.apply_district_filtering
            apply_abstract_filtering = processor_module.apply_abstract_filtering
            sheets_to_exclude = processor_module.SHEETS_TO_EXCLUDE
        else:
            from .processors import mumbai_city
            apply_district_filtering = mumbai_city.apply_district_filtering
            apply_abstract_filtering = mumbai_city.apply_abstract_filtering
            sheets_to_exclude = mumbai_city.SHEETS_TO_EXCLUDE
        
        if only_sheet is not None:
            sheet_name = SHEET_NAMES.get(only_sheet)
            if sheet_name and sheet_name in wb.sheetnames:
                source_sheet = wb[sheet_name]
                apply_district_filtering(source_sheet, only_sheet)
                wb = _copy_sheet_to_new_workbook(source_sheet, sheet_name)
        else:
            # Remove excluded sheets
            sheets_to_remove = [
                name for name in wb.sheetnames if name in sheets_to_exclude
            ]
            for sheet_name in sheets_to_remove:
                wb.remove(wb[sheet_name])
            
            # Apply filtering to remaining sheets
            for sheet_key, sheet_name in SHEET_NAMES.items():
                if sheet_name in wb.sheetnames:
                    source_sheet = wb[sheet_name]
                    apply_district_filtering(source_sheet, sheet_key)
            
            # Apply abstract filtering
            if "Distrs.wise Abstract" in wb.sheetnames:
                abstract_sheet = wb["Distrs.wise Abstract"]
                apply_abstract_filtering(abstract_sheet)
    elif only_sheet is not None:
        sheet_name = SHEET_NAMES.get(only_sheet)
        if sheet_name and sheet_name in wb.sheetnames:
            source_sheet = wb[sheet_name]
            wb = _copy_sheet_to_new_workbook(source_sheet, sheet_name)

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def export_original_workbook_async(
    db: Session,
    only_sheet: Optional[str] = None,
    user_district: Optional[str] = None,
    sub_scheme_code: Optional[str] = None,
    fiscal_year: Optional[str] = None
) -> StreamingResponse:
    """
    Export original workbook with production-grade throttling.
    
    THIS IS THE RECOMMENDED METHOD FOR PRODUCTION USE.
    
    Features:
    - Limits concurrent exports to prevent memory exhaustion
    - 60-second timeout for long-running exports
    - Proper error handling with 503/504 status codes
    - Consistent cache-prevention headers
    
    Args:
        db: Database session for data queries
        only_sheet: If specified, export only this sheet
        user_district: User's district for row filtering
        sub_scheme_code: Sub-scheme code for template selection
        fiscal_year: Fiscal year for data population
    
    Returns:
        StreamingResponse with throttling protection
    
    Raises:
        HTTPException 503: Server busy (too many concurrent exports)
        HTTPException 504: Export timeout
        HTTPException 500: Template loading or generation failure
    """
    base_filename = "original_format" if only_sheet is None else f"{only_sheet}_original_format"
    
    def generate():
        try:
            return _generate_workbook(db, only_sheet, user_district, sub_scheme_code, fiscal_year)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Failed to generate Excel: {e}")
    
    return await ExcelExportService.export_with_throttle(
        export_fn=generate,
        filename=base_filename,
        fiscal_year=fiscal_year
    )


def export_original_workbook(
    db: Session,
    only_sheet: Optional[str] = None,
    user_district: Optional[str] = None,
    sub_scheme_code: Optional[str] = None,
    fiscal_year: Optional[str] = None
) -> StreamingResponse:
    """
    Export original workbook template with populated data (synchronous).
    
    NOTE: For high-traffic production endpoints, use export_original_workbook_async
    which includes throttling protection.
    
    Args:
        db: Database session for data queries.
        only_sheet: If specified, export only this sheet (e.g., "post_status").
        user_district: User's district for row filtering.
        sub_scheme_code: Sub-scheme code for template selection.
        fiscal_year: Fiscal year for data population.
    
    Returns:
        StreamingResponse with cache-prevention headers.
    
    Raises:
        HTTPException: If template loading or Excel generation fails.
    """
    try:
        output = _generate_workbook(db, only_sheet, user_district, sub_scheme_code, fiscal_year)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate Excel: {e}")

    base_filename = "original_format" if only_sheet is None else f"{only_sheet}_original_format"
    return ExcelExportService.create_response(
        content=output,
        base_filename=base_filename,
        fiscal_year=fiscal_year
    )


