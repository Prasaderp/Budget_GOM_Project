from typing import Optional, List, Dict
from collections import defaultdict
from fastapi import HTTPException
from sqlalchemy.orm import Session
from openpyxl.workbook import Workbook
from openpyxl.utils import get_column_letter

from src.utils_scheme import get_scheme_models
from src.utils_da_rate import get_da_rate
from ...config import HRA_RATE_MAP, SCHEME_DISTRICTS

# ==========================================
# 1. THE TRUTH TABLE (MAPPING)
# Map your DB Designation Name (Key) to the Exact Excel Row Number (Value)
# ==========================================

DCO_PERMANENT_MAP = {
    # Class 1 & 2
    "Divisional Commissioner": 7,
    "Additional Commissioner": 8,
    "Tehsildar": 9,
    "Naib Tehsildar": 10,
    
    # Class 3 (Adjust these row numbers based on your Excel!)
    "Stenographer (Higher)": 12,
    "Head Clerk (Awwal Karkun)": 13,
    "Clerk-Typist": 14,
    
    # Class 4
    "Peon/Naik/Havaldar": 16, # Example, check your sheet if this is row 16 or 19
}

DCO_TEMPORARY_MAP = {
    # Class 1 & 2 (Starts around row 26 usually)
    "Deputy Commissioner": 26,
    "Deputy Collector": 27,
    "Assistant Director Town Planning": 28,
    "Assistant Director": 29,
    "Naib Tehsildar": 30,
    "Naib Tehsildar (Ulhasnagar)": 31,
    "Accounts Officer": 32,
    "Planning Assistant": 33,
    "Assistant Accounts Officer": 34,
    "Law Officer (Honorarium)": 35, # Skip row 36 if it's Total
    
    # Class 3
    "Sub-Accountant": 37,
    "Stenographer (Selection Grade)": 38,
    "Stenographer (Lower)": 39,
    "Stenographer (U.V. Ulhasnagar)": 40,
    "Junior Typist": 41,
    "Cashier/Senior Pay Scale/Recovery Agent/Rent Collector": 42,
    "Clerk-Typist": 43,
    "Vehicle Driver": 44,
    "Sanitation Inspector": 45, # Skip 46 if total
    
    # Class 4
    "Peon/Naik/Havaldar": 47,
    "Watchman/Guard": 48,
    "Lift Operator": 49,
    "Porter": 50,
    "Sweeper/Cleaner": 51,
    "Worker": 52,
}

# ==========================================
# 2. THE WRITER CODE
# ==========================================

def _write(ws, cell_addr: Optional[str], value):
    """Write value to cell if address is valid."""
    if cell_addr:
        # We don't need MergedCell checks anymore because 
        # we are mapping to known valid rows only.
        ws[cell_addr].value = value if value is not None else None


def populate_budget_post_details(
    wb: Workbook,
    db: Session,
    sub_scheme_code: Optional[str] = None,
    fiscal_year: Optional[str] = None
):
    from ...config import SHEET_NAMES

    sheet_name = SHEET_NAMES.get("budget_post_details")
    if sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=500, detail=f"Sheet '{sheet_name}' not found")
    ws = wb[sheet_name]

    da_rate = get_da_rate(db, fiscal_year)
    BudgetPostDetails, _, _, _ = get_scheme_models(sub_scheme_code)

    # We process by Category using the specific maps
    _write_mapped_block(ws, db, BudgetPostDetails, "DCO Staff", "Permanent", DCO_PERMANENT_MAP, da_rate, fiscal_year)
    _write_mapped_block(ws, db, BudgetPostDetails, "DCO Staff", "Temporary", DCO_TEMPORARY_MAP, da_rate, fiscal_year)


def _write_mapped_block(
    ws,
    db: Session,
    model,
    dco_unit: str,
    category: str,
    row_map: Dict[str, int],
    da_rate: float,
    fiscal_year: Optional[str]
):
    """
    Writes data to exact rows defined in the mapping dictionary.
    No guessing, no skipping loops. Pure coordinate writing.
    """
    # 1. Fetch Records
    query = db.query(model).filter(model.district == dco_unit, model.category == category)
    if fiscal_year:
        query = query.filter(model.fiscal_year == fiscal_year)
    records = query.all()

    # 2. Aggregate Data
    # We pass the row_map keys effectively as the 'valid designations'
    agg = _aggregate_records(records, row_map, da_rate)

    # 3. Write Data
    _write_aggregated_data(ws, row_map, agg)


def _aggregate_records(
    records: List,
    row_map: Dict[str, int],
    da_rate: float
) -> Dict[str, Dict[str, int]]:
    agg = defaultdict(lambda: {
        "sanctioned_posts_prev1": 0, "sanctioned_posts_curr": 0,
        "special_pay": 0, "basic_pay": 0, "grade_pay": 0,
        "dearness_allowance": 0, "local_supplementary_allowance": 0,
        "house_rent_allowance": 0, "vehicle_allowance": 0,
        "washing_allowance": 0, "cash_allowance": 0, "footwear_allowance_other": 0,
    })

    for rec in records:
        desig = rec.designation
        
        # KEY FIX: If the DB designation isn't in our map, we skip it (or log it)
        if desig not in row_map:
            print(f"WARNING: Designation '{desig}' found in DB but not in Row Map. Skipping.")
            continue

        basic_pay = int(float(rec.basic_pay or 0))
        grade_pay = int(rec.grade_pay or 0)
        base_salary = basic_pay + grade_pay
        hra_rate_val = HRA_RATE_MAP.get(rec.hra_rate, 0.3)

        agg[desig]["sanctioned_posts_prev1"] += int(rec.sanctioned_posts_prev1 or 0)
        agg[desig]["sanctioned_posts_curr"] += int(rec.sanctioned_posts_curr or 0)
        agg[desig]["special_pay"] += int(rec.special_pay or 0)
        agg[desig]["basic_pay"] += basic_pay
        agg[desig]["grade_pay"] += grade_pay
        agg[desig]["dearness_allowance"] += int(base_salary * da_rate)
        agg[desig]["local_supplementary_allowance"] += int(rec.local_supplementary_allowance or 0)
        agg[desig]["house_rent_allowance"] += int(base_salary * hra_rate_val)
        agg[desig]["vehicle_allowance"] += int(rec.vehicle_allowance or 0)
        agg[desig]["washing_allowance"] += int(rec.washing_allowance or 0)
        agg[desig]["cash_allowance"] += int(rec.cash_allowance or 0)
        agg[desig]["footwear_allowance_other"] += int(rec.footwear_allowance_other or 0)

    return agg


def _write_aggregated_data(ws, row_map: Dict[str, int], agg: Dict):
    col_map = [
        ("D", "sanctioned_posts_prev1"), ("E", "sanctioned_posts_curr"),
        ("F", "special_pay"), ("G", "basic_pay"), ("H", "grade_pay"),
        ("J", "dearness_allowance"), ("K", "local_supplementary_allowance"),
        ("L", "house_rent_allowance"), ("M", "vehicle_allowance"),
        ("N", "washing_allowance"), ("O", "cash_allowance"), ("P", "footwear_allowance_other"),
    ]

    for desig, data in agg.items():
        # Look up the row number from the map
        row = row_map.get(desig)
        if not row: 
            continue
            
        for col, field in col_map:
            _write(ws, f"{col}{row}", data[field])

# ... Include your DCO_ROW_RANGES and filtering logic below as before ...