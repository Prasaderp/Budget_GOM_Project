"""DCO Staff processor for sub-scheme 20530387.

This scheme has NO districts - only DCO units.
"""
from openpyxl.utils import get_column_letter

# Row ranges for DCO Staff filtering
# NOTE: Adjust these values after verifying with actual template
DCO_ROW_RANGES = {
    "budget_post_details": (1, 54),
    "post_status": (1, 32),
    "post_expenses": (1, 16),
    "unit_expenditure": (1, 22)
}

SHEETS_TO_EXCLUDE = []


def apply_dco_filtering(source_sheet, sheet_key: str):
    """Filter rows for DCO Staff."""
    range_info = DCO_ROW_RANGES.get(sheet_key)
    if not range_info:
        return

    start_row, end_row = range_info[:2]
    cols_to_include = range_info[2] if len(range_info) > 2 else None

    rows_to_delete = [r for r in range(1, source_sheet.max_row + 1)
                      if not (start_row <= r <= end_row)]
    for row in reversed(rows_to_delete):
        source_sheet.delete_rows(row)

    if cols_to_include:
        cols_to_delete = [c for c in range(1, source_sheet.max_column + 1)
                          if get_column_letter(c) not in cols_to_include]
        for col in reversed(cols_to_delete):
            source_sheet.delete_cols(col)


# Alias for compatibility
apply_district_filtering = apply_dco_filtering
