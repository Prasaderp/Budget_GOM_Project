import io
import os
from typing import Optional
from fastapi import HTTPException
from starlette.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import load_workbook, Workbook
from copy import copy

from ..config import EXCEL_TEMPLATE_PATH, SHEET_NAMES, DCO_STAFF_IDENTIFIER
from .populators.budget_post_details import populate_budget_post_details
from .populators.post_status import populate_post_status
from .populators.post_expenses import populate_post_expenses
from .populators.unit_expenditure import populate_unit_expenditure


def _get_template_path(sub_scheme_code: Optional[str] = None) -> str:
    if sub_scheme_code:
        parent_scheme = sub_scheme_code[:4]
        template_path = f"excel_templates/s{parent_scheme}/subs/s{sub_scheme_code}/original_template.xlsx"
        if os.path.exists(template_path):
            return template_path
    return EXCEL_TEMPLATE_PATH


def _get_processor_module(district: str):
    DISTRICT_PROCESSORS = {
        "Mumbai City": "mumbai_city",
        "Mumbai Suburban": "mumbai_suburban",
        "Thane": "thane",
        "Palghar": "palghar",
        "Raigad": "raigad",
        "Ratnagiri": "ratnagiri",
        "Sindhudurg": "sindhudurg",
        DCO_STAFF_IDENTIFIER: "dco_staff"
    }
    
    processor_name = DISTRICT_PROCESSORS.get(district)
    if not processor_name:
        return None
    
    try:
        processor_module = __import__(
            f'src.schemes.s2053.subs.s20530028.excel_export.processors.{processor_name}',
            fromlist=[processor_name]
        )
        return processor_module
    except ImportError:
        from .processors import mumbai_city
        return mumbai_city


def _copy_sheet_to_new_workbook(source_sheet, sheet_name: str) -> Workbook:
    new_wb = Workbook()
    new_wb.remove(new_wb.active)
    target_sheet = new_wb.create_sheet(sheet_name)
    
    for row in source_sheet.iter_rows():
        for cell in row:
            if hasattr(cell, 'coordinate'):
                target_cell = target_sheet[cell.coordinate]
                target_cell.value = cell.value
                if hasattr(cell, 'has_style') and cell.has_style:
                    target_cell.font = copy(cell.font)
                    target_cell.border = copy(cell.border)
                    target_cell.fill = copy(cell.fill)
                    target_cell.number_format = cell.number_format
                    target_cell.protection = copy(cell.protection)
                    target_cell.alignment = copy(cell.alignment)
    
    for merged_range in source_sheet.merged_cells.ranges:
        target_sheet.merge_cells(str(merged_range))
    
    for col_letter, col_dim in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[col_letter].width = col_dim.width
    
    for row_num, row_dim in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[row_num].height = row_dim.height
    
    return new_wb


def export_original_workbook(
    db: Session,
    only_sheet: Optional[str] = None,
    user_district: Optional[str] = None,
    sub_scheme_code: Optional[str] = None,
    fiscal_year: Optional[str] = None
) -> StreamingResponse:
    template_path = _get_template_path(sub_scheme_code)
    try:
        wb = load_workbook(template_path, data_only=False)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not load Excel template: {e}")

    try:
        if only_sheet in (None, "budget_post_details"):
            populate_budget_post_details(wb, db, sub_scheme_code, fiscal_year)
        if only_sheet in (None, "post_status"):
            populate_post_status(wb, db, sub_scheme_code, fiscal_year)
        if only_sheet in (None, "post_expenses"):
            populate_post_expenses(wb, db, sub_scheme_code, fiscal_year)
        if only_sheet in (None, "unit_expenditure"):
            populate_unit_expenditure(wb, db, sub_scheme_code, fiscal_year)

        if user_district is not None:
            processor_module = _get_processor_module(user_district)
            if processor_module:
                apply_district_filtering = processor_module.apply_district_filtering
                apply_abstract_filtering = processor_module.apply_abstract_filtering
                sheets_to_exclude = processor_module.SHEETS_TO_EXCLUDE
            else:
                from .processors import mumbai_city
                apply_district_filtering = mumbai_city.apply_district_filtering
                apply_abstract_filtering = mumbai_city.apply_abstract_filtering
                sheets_to_exclude = mumbai_city.SHEETS_TO_EXCLUDE
            
            if only_sheet is not None:
                sheet_name = SHEET_NAMES.get(only_sheet)
                if sheet_name and sheet_name in wb.sheetnames:
                    source_sheet = wb[sheet_name]
                    apply_district_filtering(source_sheet, only_sheet)
                    wb = _copy_sheet_to_new_workbook(source_sheet, sheet_name)
            else:
                sheets_to_remove = []
                for sheet_name in wb.sheetnames:
                    if sheet_name in sheets_to_exclude:
                        sheets_to_remove.append(sheet_name)
                
                for sheet_name in sheets_to_remove:
                    wb.remove(wb[sheet_name])
                
                for sheet_key, sheet_name in SHEET_NAMES.items():
                    if sheet_name in wb.sheetnames:
                        source_sheet = wb[sheet_name]
                        apply_district_filtering(source_sheet, sheet_key)
                
                if "Distrs.wise Abstract" in wb.sheetnames:
                    abstract_sheet = wb["Distrs.wise Abstract"]
                    apply_abstract_filtering(abstract_sheet)
        elif only_sheet is not None:
            sheet_name = SHEET_NAMES.get(only_sheet)
            if sheet_name and sheet_name in wb.sheetnames:
                source_sheet = wb[sheet_name]
                wb = _copy_sheet_to_new_workbook(source_sheet, sheet_name)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = "original_format.xlsx" if only_sheet is None else f"{only_sheet}_original_format.xlsx"
        headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
        return StreamingResponse(output, headers=headers, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate Excel: {e}")
