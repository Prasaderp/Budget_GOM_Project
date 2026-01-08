from openpyxl.utils import get_column_letter

DISTRICT_ROW_RANGES = {
    "budget_post_details": (142, 188),
    "post_status": (99, 132),
    "post_expenses": (49, 64),
    "unit_expenditure": (70, 92)
}

DISTRICT_WISE_ABSTRACT_RANGE = (1, 22, ['A', 'B', 'F'])

SHEETS_TO_EXCLUDE = ["Page-5", "Table-D"]

def apply_district_filtering(source_sheet, sheet_key):
    range_info = DISTRICT_ROW_RANGES.get(sheet_key)
    
    if range_info:
        if len(range_info) == 3:
            start_row, end_row, cols_to_include = range_info
        else:
            start_row, end_row = range_info
            cols_to_include = None
        
        from copy import copy
        import re
        
        extracted_data = []
        row_heights = {}
        merged_ranges = []
        
        for row_num in range(start_row, end_row + 1):
            row_data = []
            if row_num in source_sheet.row_dimensions:
                row_heights[row_num - start_row + 1] = source_sheet.row_dimensions[row_num].height
            
            for col_num in range(1, source_sheet.max_column + 1):
                cell = source_sheet.cell(row=row_num, column=col_num)
                
                cell_value = cell.value
                if cell_value and isinstance(cell_value, str) and cell_value.startswith('='):
                    cell_value = re.sub(r'\$?([A-Z]+)\$?(\d+)', 
                                      lambda m: f"${m.group(1)}${int(m.group(2)) - start_row + 1}", 
                                      cell_value)
                
                row_data.append({
                    'value': cell_value,
                    'font': copy(cell.font) if hasattr(cell, 'font') else None,
                    'border': copy(cell.border) if hasattr(cell, 'border') else None,
                    'fill': copy(cell.fill) if hasattr(cell, 'fill') else None,
                    'number_format': cell.number_format,
                    'protection': copy(cell.protection) if hasattr(cell, 'protection') else None,
                    'alignment': copy(cell.alignment) if hasattr(cell, 'alignment') else None
                })
            extracted_data.append(row_data)
        
        for merged_range in source_sheet.merged_cells.ranges:
            min_row, min_col, max_row, max_col = merged_range.bounds
            if start_row <= min_row <= end_row and start_row <= max_row <= end_row:
                new_min_row = min_row - start_row + 1
                new_max_row = max_row - start_row + 1
                merged_ranges.append((new_min_row, min_col, new_max_row, max_col))
        
        for row_num in range(source_sheet.max_row, 0, -1):
            source_sheet.delete_rows(row_num)
        
        for idx, row_data in enumerate(extracted_data):
            target_row = idx + 1
            if target_row in row_heights:
                source_sheet.row_dimensions[target_row].height = row_heights[target_row]
            
            for col_idx, cell_data in enumerate(row_data):
                cell = source_sheet.cell(row=target_row, column=col_idx + 1)
                cell.value = cell_data['value']
                if cell_data['font']:
                    cell.font = cell_data['font']
                if cell_data['border']:
                    cell.border = cell_data['border']
                if cell_data['fill']:
                    cell.fill = cell_data['fill']
                if cell_data['number_format']:
                    cell.number_format = cell_data['number_format']
                if cell_data['protection']:
                    cell.protection = cell_data['protection']
                if cell_data['alignment']:
                    cell.alignment = cell_data['alignment']
        
        for new_min_row, min_col, new_max_row, max_col in merged_ranges:
            source_sheet.merge_cells(start_row=new_min_row, start_column=min_col, end_row=new_max_row, end_column=max_col)
        
        if cols_to_include:
            cols_to_delete = []
            for col_num in range(1, source_sheet.max_column + 1):
                col_letter = get_column_letter(col_num)
                if col_letter not in cols_to_include:
                    cols_to_delete.append(col_num)
            
            for col_num in reversed(cols_to_delete):
                source_sheet.delete_cols(col_num)

def apply_abstract_filtering(source_sheet):
    start_row, end_row, cols_to_include = DISTRICT_WISE_ABSTRACT_RANGE
    
    for row_num in range(5, 20):
        for col_num in range(1, source_sheet.max_column + 1):
            col_letter = get_column_letter(col_num)
            if col_letter == 'F':
                cell = source_sheet.cell(row=row_num, column=col_num)
                if cell.value and isinstance(cell.value, str) and "'Page-4'!" in cell.value:
                    new_row_ref = row_num - 5 + 7
                    cell.value = f"='Page-4'!F{new_row_ref}"
    
    total_row = 20
    for col_num in range(1, source_sheet.max_column + 1):
        col_letter = get_column_letter(col_num)
        if col_letter == 'F':
            cell = source_sheet.cell(row=total_row, column=col_num)
            cell.value = "=C5+C6+C7+C8+C9+C10+C11+C12+C13+C14+C15+C16+C17+C18+C19"
    
    rows_to_delete = []
    for row_num in range(1, source_sheet.max_row + 1):
        if not (start_row <= row_num <= end_row):
            rows_to_delete.append(row_num)
    
    for row_num in reversed(rows_to_delete):
        source_sheet.delete_rows(row_num)
    
    cols_to_delete = []
    for col_num in range(1, source_sheet.max_column + 1):
        col_letter = get_column_letter(col_num)
        if col_letter not in cols_to_include:
            cols_to_delete.append(col_num)
    
    for col_num in reversed(cols_to_delete):
        source_sheet.delete_cols(col_num)
