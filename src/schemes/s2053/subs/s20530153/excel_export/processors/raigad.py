from openpyxl.utils import get_column_letter

# District-specific row ranges for 20530153 - Raigad
DISTRICT_ROW_RANGES = {
    "budget_post_details": (117, 145),
    "post_status": (133, 165),
    "post_expenses": (65, 80),
    "unit_expenditure": (89, 110)
}

DISTRICT_WISE_ABSTRACT_RANGE = (1, 20, ['A', 'B', 'C'])

SHEETS_TO_EXCLUDE = ["Page-5", "Table-D"]

def apply_district_filtering(source_sheet, sheet_key):
    range_info = DISTRICT_ROW_RANGES.get(sheet_key)
    
    if range_info:
        if len(range_info) == 3:
            start_row, end_row, cols_to_include = range_info
        else:
            start_row, end_row = range_info
            cols_to_include = None
        
        rows_to_delete = []
        for row_num in range(1, source_sheet.max_row + 1):
            if not (start_row <= row_num <= end_row):
                rows_to_delete.append(row_num)
        
        for row_num in reversed(rows_to_delete):
            source_sheet.delete_rows(row_num)
        
        if cols_to_include:
            cols_to_delete = []
            for col_num in range(1, source_sheet.max_column + 1):
                col_letter = get_column_letter(col_num)
                if col_letter not in cols_to_include:
                    cols_to_delete.append(col_num)
            
            for col_num in reversed(cols_to_delete):
                source_sheet.delete_cols(col_num)

def apply_abstract_filtering(source_sheet):
    start_row, end_row, cols_to_include = DISTRICT_WISE_ABSTRACT_RANGE
    
    total_row = 20
    for col_num in range(1, source_sheet.max_column + 1):
        col_letter = get_column_letter(col_num)
        if col_letter == 'C':
            cell = source_sheet.cell(row=total_row, column=col_num)
            cell.value = "=C5+C6+C7+C8+C9+C10+C11+C12+C13+C14+C15+C16+C17+C18+C19"
    
    rows_to_delete = []
    for row_num in range(1, source_sheet.max_row + 1):
        if not (start_row <= row_num <= end_row):
            rows_to_delete.append(row_num)
    
    for row_num in reversed(rows_to_delete):
        source_sheet.delete_rows(row_num)
    
    cols_to_delete = []
    for col_num in range(1, source_sheet.max_column + 1):
        col_letter = get_column_letter(col_num)
        if col_letter not in cols_to_include:
            cols_to_delete.append(col_num)
    
    for col_num in reversed(cols_to_delete):
        source_sheet.delete_cols(col_num)

