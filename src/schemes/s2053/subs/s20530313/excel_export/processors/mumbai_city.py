"""Mumbai City district processor for sub-scheme 20530313."""
from openpyxl.utils import get_column_letter

# Row ranges for district filtering
# NOTE: Adjust these values after verifying with actual template
DISTRICT_ROW_RANGES = {
    "budget_post_details": (1, 19),
    "post_status": (1, 32),
    "post_expenses": (1, 15),
    "unit_expenditure": (1, 20)
}

DISTRICT_WISE_ABSTRACT_RANGE = (1, 20, ['A', 'B', 'C'])
SHEETS_TO_EXCLUDE = ["Page-5", "Table-D"]


def apply_district_filtering(source_sheet, sheet_key: str):
    """Filter rows for Mumbai City district."""
    range_info = DISTRICT_ROW_RANGES.get(sheet_key)
    if not range_info:
        return

    start_row, end_row = range_info[:2]
    cols_to_include = range_info[2] if len(range_info) > 2 else None

    rows_to_delete = [r for r in range(1, source_sheet.max_row + 1)
                      if not (start_row <= r <= end_row)]
    for row in reversed(rows_to_delete):
        source_sheet.delete_rows(row)

    if cols_to_include:
        cols_to_delete = [c for c in range(1, source_sheet.max_column + 1)
                          if get_column_letter(c) not in cols_to_include]
        for col in reversed(cols_to_delete):
            source_sheet.delete_cols(col)


def apply_abstract_filtering(source_sheet):
    """Filter abstract sheet for Mumbai City."""
    start_row, end_row, cols_to_include = DISTRICT_WISE_ABSTRACT_RANGE

    for col_num in range(1, source_sheet.max_column + 1):
        if get_column_letter(col_num) == 'C':
            source_sheet.cell(row=20, column=col_num).value = \
                "=C5+C6+C7+C8+C9+C10+C11+C12+C13+C14+C15+C16+C17+C18+C19"

    rows_to_delete = [r for r in range(1, source_sheet.max_row + 1)
                      if not (start_row <= r <= end_row)]
    for row in reversed(rows_to_delete):
        source_sheet.delete_rows(row)

    cols_to_delete = [c for c in range(1, source_sheet.max_column + 1)
                      if get_column_letter(c) not in cols_to_include]
    for col in reversed(cols_to_delete):
        source_sheet.delete_cols(col)
