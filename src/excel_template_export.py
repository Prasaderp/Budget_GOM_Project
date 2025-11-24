import io
from typing import Dict, Tuple, Optional, List
from fastapi import HTTPException
from starlette.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from src import models
from src.config import ORIGINAL_XLSX_PATH, ORIGINAL_SHEET_NAMES


def _write(ws, cell_addr: Optional[str], value):
    if not cell_addr:
        return
    ws[cell_addr].value = value if value is not None else None

def export_original_workbook(db: Session, only_sheet: Optional[str] = None, user_district: Optional[str] = None) -> StreamingResponse:
    try:
        wb = load_workbook(ORIGINAL_XLSX_PATH, data_only=False)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not load Excel template: {e}")

    try:
        if only_sheet in (None, "budget_post_details"):
            populate_budget_post_details(wb, db)
        if only_sheet in (None, "post_status"):
            populate_post_status(wb, db)
        if only_sheet in (None, "post_expenses"):
            populate_post_expenses(wb, db)
        if only_sheet in (None, "unit_expenditure"):
            populate_unit_expenditure(wb, db)

        if user_district is not None:
            DISTRICT_PROCESSORS = {
                "Mumbai City": "mumbai_city",
                "Mumbai Suburban": "mumbai_suburban",
                "Thane": "thane",
                "Palghar": "palghar",
                "Raigad": "raigad",
                "Ratnagiri": "ratnagiri",
                "Sindhudurg": "sindhudurg"
            }
            
            processor_name = DISTRICT_PROCESSORS.get(user_district)
            if processor_name:
                try:
                    processor_module = __import__(f'src.excel_processors.{processor_name}', fromlist=[processor_name])
                    apply_district_filtering = processor_module.apply_district_filtering
                    apply_abstract_filtering = processor_module.apply_abstract_filtering
                    sheets_to_exclude = processor_module.SHEETS_TO_EXCLUDE
                except ImportError:
                    from src.excel_processors.mumbai_city import apply_district_filtering, apply_abstract_filtering, SHEETS_TO_EXCLUDE as sheets_to_exclude
            
            if only_sheet is not None:
                from openpyxl import Workbook
                from copy import copy
                
                sheet_name = ORIGINAL_SHEET_NAMES.get(only_sheet)
                if sheet_name and sheet_name in wb.sheetnames:
                    source_sheet = wb[sheet_name]
                    apply_district_filtering(source_sheet, only_sheet)
                    
                    new_wb = Workbook()
                    new_wb.remove(new_wb.active)
                    target_sheet = new_wb.create_sheet(sheet_name)
                    
                    for row in source_sheet.iter_rows():
                        for cell in row:
                            if hasattr(cell, 'coordinate'):
                                target_cell = target_sheet[cell.coordinate]
                                target_cell.value = cell.value
                                if hasattr(cell, 'has_style') and cell.has_style:
                                    target_cell.font = copy(cell.font)
                                    target_cell.border = copy(cell.border)
                                    target_cell.fill = copy(cell.fill)
                                    target_cell.number_format = cell.number_format
                                    target_cell.protection = copy(cell.protection)
                                    target_cell.alignment = copy(cell.alignment)
                    
                    for merged_range in source_sheet.merged_cells.ranges:
                        target_sheet.merge_cells(str(merged_range))
                    
                    for col_letter, col_dim in source_sheet.column_dimensions.items():
                        target_sheet.column_dimensions[col_letter].width = col_dim.width
                    
                    for row_num, row_dim in source_sheet.row_dimensions.items():
                        target_sheet.row_dimensions[row_num].height = row_dim.height
                    
                    wb = new_wb
            else:
                sheets_to_remove = []
                for sheet_name in wb.sheetnames:
                    if sheet_name in sheets_to_exclude:
                        sheets_to_remove.append(sheet_name)
                
                for sheet_name in sheets_to_remove:
                    wb.remove(wb[sheet_name])
                
                for sheet_key, sheet_name in ORIGINAL_SHEET_NAMES.items():
                    if sheet_name in wb.sheetnames:
                        source_sheet = wb[sheet_name]
                        apply_district_filtering(source_sheet, sheet_key)
                
                if "Distrs.wise Abstract" in wb.sheetnames:
                    abstract_sheet = wb["Distrs.wise Abstract"]
                    apply_abstract_filtering(abstract_sheet)
        elif only_sheet is not None:
            from openpyxl import Workbook
            from copy import copy
            new_wb = Workbook()
            new_wb.remove(new_wb.active)
            
            sheet_name = ORIGINAL_SHEET_NAMES.get(only_sheet)
            if sheet_name and sheet_name in wb.sheetnames:
                source_sheet = wb[sheet_name]
                target_sheet = new_wb.create_sheet(sheet_name)
                
                for row in source_sheet.iter_rows():
                    for cell in row:
                        if hasattr(cell, 'coordinate'):
                            target_cell = target_sheet[cell.coordinate]
                            target_cell.value = cell.value
                            if hasattr(cell, 'has_style') and cell.has_style:
                                target_cell.font = copy(cell.font)
                                target_cell.border = copy(cell.border)
                                target_cell.fill = copy(cell.fill)
                                target_cell.number_format = cell.number_format
                                target_cell.protection = copy(cell.protection)
                                target_cell.alignment = copy(cell.alignment)
                
                for merged_range in source_sheet.merged_cells.ranges:
                    target_sheet.merge_cells(str(merged_range))
                
                for col_letter, col_dim in source_sheet.column_dimensions.items():
                    target_sheet.column_dimensions[col_letter].width = col_dim.width
                
                for row_num, row_dim in source_sheet.row_dimensions.items():
                    target_sheet.row_dimensions[row_num].height = row_dim.height
                
                wb = new_wb

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = "original_format.xlsx" if only_sheet is None else f"{only_sheet}_original_format.xlsx"
        headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
        return StreamingResponse(output, headers=headers, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate Excel: {e}")

def populate_budget_post_details(wb, db: Session):
    sheet_name = ORIGINAL_SHEET_NAMES.get("budget_post_details")
    if sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=500, detail=f"Sheet '{sheet_name}' not found in original workbook")
    ws = wb[sheet_name]

    def addr(col: str, row: int) -> str:
        return f"{col}{row}"

    from collections import defaultdict
    from config import DESIGNATIONS

    perm_designations = DESIGNATIONS[:14]
    temp_designations = [
        'Collector',
        'Additional Collector',
        'Deputy Collector / Probationary Deputy Collector',
        'Tehsildar/Additional Tehsildar/Chitnis (Clerk/Secretary)/Probationary Tehsildar',
        'Naib Tehsildar/Probationary Naib Tehsildar',
        'Accounts Officer',
        'Asst. Accounts Officer',
        'Law Officer (Honorarium)',
        'Deputy Accountant',
        'Head Clerk/Deputy Accountant',
        'Circle Officer',
        'Clerk/Land Surveyor/Recovery Clerk',
        'Stenographer (Higher)',
        'Stenographer (Lower)/Probationary Land Surveyor/Draftsman/Shirastedar',
        'Vehicle Driver',
        'Telephone Operator/Steno-Typist(Law Officer Asst.)',
        'Peon/Naik/Havaldar/Watchman/Cleaner',
    ]

    def write_block(district: str, category: str, designations: List[str], start_row: int):
        row_map = {desig: row for desig, row in zip(designations, range(start_row, start_row + len(designations)))}
        records: List[models.BudgetPostDetails] = (
            db.query(models.BudgetPostDetails)
            .filter(models.BudgetPostDetails.district == district, models.BudgetPostDetails.category == category)
            .all()
        )
        agg = defaultdict(lambda: {
            "sanctioned_posts_2024_25": 0,
            "sanctioned_posts_2025_26": 0,
            "special_pay": 0,
            "basic_pay": 0,
            "grade_pay": 0,
            "local_supplementary_allowance": 0,
            "vehicle_allowance": 0,
            "washing_allowance": 0,
            "cash_allowance": 0,
            "footwear_allowance_other": 0,
        })
        for it in records:
            key = it.designation
            if key not in row_map:
                continue
            agg[key]["sanctioned_posts_2024_25"] += int(it.sanctioned_posts_2024_25 or 0)
            agg[key]["sanctioned_posts_2025_26"] += int(it.sanctioned_posts_2025_26 or 0)
            agg[key]["special_pay"] += int(it.special_pay or 0)
            agg[key]["basic_pay"] += int(it.basic_pay or 0)
            agg[key]["grade_pay"] += int(it.grade_pay or 0)
            agg[key]["local_supplementary_allowance"] += int(it.local_supplementary_allowance or 0)
            agg[key]["vehicle_allowance"] += int(it.vehicle_allowance or 0)
            agg[key]["washing_allowance"] += int(it.washing_allowance or 0)
            agg[key]["cash_allowance"] += int(it.cash_allowance or 0)
            agg[key]["footwear_allowance_other"] += int(it.footwear_allowance_other or 0)
        for desig, row in row_map.items():
            vals = agg.get(desig)
            if not vals:
                continue
            _write(ws, addr("D", row), vals["sanctioned_posts_2024_25"])  # 2024-25
            _write(ws, addr("E", row), vals["sanctioned_posts_2025_26"])  # 2025-26
            _write(ws, addr("F", row), vals["special_pay"])               # Special Pay
            _write(ws, addr("G", row), vals["basic_pay"])                 # Basic Pay
            _write(ws, addr("H", row), vals["grade_pay"])                 # Grade Pay
            _write(ws, addr("K", row), vals["local_supplementary_allowance"])  # Local Supp. Allowance
            _write(ws, addr("M", row), vals["vehicle_allowance"])         # Vehicle Allowance
            _write(ws, addr("N", row), vals["washing_allowance"])         # Washing Allowance
            _write(ws, addr("O", row), vals["cash_allowance"])            # Cash Allowance
            _write(ws, addr("P", row), vals["footwear_allowance_other"])  # Footwear/Other

    blocks = [
        ("Mumbai City", 7, 29),
        ("Mumbai Suburban", 54, 76),
        ("Thane", 101, 123),
        ("Palghar", 148, 170),
        ("Raigad", 195, 217),
        ("Ratnagiri", 242, 264),
        ("Sindhudurg", 289, 311),
    ]

    for district, perm_start, temp_start in blocks:
        write_block(district, "Permanent", perm_designations, perm_start)
        write_block(district, "Temporary", temp_designations, temp_start)

def populate_post_status(wb, db: Session):
    sheet_name = ORIGINAL_SHEET_NAMES.get("post_status")
    if sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=500, detail=f"Sheet '{sheet_name}' not found in original workbook")
    ws = wb[sheet_name]
    def addr(col: str, row: int) -> str:
        return f"{col}{row}"
    filled_cols = {"Class-1 & 2": "C", "Class-3": "D", "Class-4": "E"}
    vacant_cols = {"Class-1 & 2": "G", "Class-3": "H", "Class-4": "I"}

    perm_records: List[models.PostStatus] = (
        db.query(models.PostStatus)
        .filter(models.PostStatus.district == "Mumbai City", models.PostStatus.category == "Permanent")
        .all()
    )
    for r in perm_records:
        col = filled_cols.get(r.class_type) if r.status == "Filled" else vacant_cols.get(r.class_type) if r.status == "Vacant" else None
        if not col:
            continue
        _write(ws, addr(col, 7), r.posts)
        _write(ws, addr(col, 8), r.salary)
        _write(ws, addr(col, 9), r.grade_pay)
        _write(ws, addr(col, 11), r.special_pay)
        _write(ws, addr(col, 12), r.dearness_allowance)
        _write(ws, addr(col, 13), r.local_supplementary_allowance)
        _write(ws, addr(col, 14), r.house_rent_allowance)
        _write(ws, addr(col, 15), r.travel_allowance)
        _write(ws, addr(col, 16), r.other)

    temp_records: List[models.PostStatus] = (
        db.query(models.PostStatus)
        .filter(models.PostStatus.district == "Mumbai City", models.PostStatus.category == "Temporary")
        .all()
    )
    for r in temp_records:
        col = filled_cols.get(r.class_type) if r.status == "Filled" else vacant_cols.get(r.class_type) if r.status == "Vacant" else None
        if not col:
            continue
        _write(ws, addr(col, 22), r.posts)
        _write(ws, addr(col, 23), r.salary)
        _write(ws, addr(col, 24), r.grade_pay)
        _write(ws, addr(col, 26), r.special_pay)
        _write(ws, addr(col, 27), r.dearness_allowance)
        _write(ws, addr(col, 28), r.local_supplementary_allowance)
        _write(ws, addr(col, 29), r.house_rent_allowance)
        _write(ws, addr(col, 30), r.travel_allowance)
        _write(ws, addr(col, 31), r.other)

    ms_perm_records: List[models.PostStatus] = (
        db.query(models.PostStatus)
        .filter(models.PostStatus.district == "Mumbai Suburban", models.PostStatus.category == "Permanent")
        .all()
    )
    for r in ms_perm_records:
        col = filled_cols.get(r.class_type) if r.status == "Filled" else vacant_cols.get(r.class_type) if r.status == "Vacant" else None
        if not col:
            continue
        _write(ws, addr(col, 40), r.posts)
        _write(ws, addr(col, 41), r.salary)
        _write(ws, addr(col, 42), r.grade_pay)
        _write(ws, addr(col, 44), r.special_pay)
        _write(ws, addr(col, 45), r.dearness_allowance)
        _write(ws, addr(col, 46), r.local_supplementary_allowance)
        _write(ws, addr(col, 47), r.house_rent_allowance)
        _write(ws, addr(col, 48), r.travel_allowance)
        _write(ws, addr(col, 49), r.other)

    ms_temp_records: List[models.PostStatus] = (
        db.query(models.PostStatus)
        .filter(models.PostStatus.district == "Mumbai Suburban", models.PostStatus.category == "Temporary")
        .all()
    )
    for r in ms_temp_records:
        col = filled_cols.get(r.class_type) if r.status == "Filled" else vacant_cols.get(r.class_type) if r.status == "Vacant" else None
        if not col:
            continue
        _write(ws, addr(col, 55), r.posts)
        _write(ws, addr(col, 56), r.salary)
        _write(ws, addr(col, 57), r.grade_pay)
        _write(ws, addr(col, 59), r.special_pay)
        _write(ws, addr(col, 60), r.dearness_allowance)
        _write(ws, addr(col, 61), r.local_supplementary_allowance)
        _write(ws, addr(col, 62), r.house_rent_allowance)
        _write(ws, addr(col, 63), r.travel_allowance)
        _write(ws, addr(col, 64), r.other)

    tn_perm_records: List[models.PostStatus] = (
        db.query(models.PostStatus)
        .filter(models.PostStatus.district == "Thane", models.PostStatus.category == "Permanent")
        .all()
    )
    for r in tn_perm_records:
        col = filled_cols.get(r.class_type) if r.status == "Filled" else vacant_cols.get(r.class_type) if r.status == "Vacant" else None
        if not col:
            continue
        _write(ws, addr(col, 73), r.posts)
        _write(ws, addr(col, 74), r.salary)
        _write(ws, addr(col, 75), r.grade_pay)
        _write(ws, addr(col, 77), r.special_pay)
        _write(ws, addr(col, 78), r.dearness_allowance)
        _write(ws, addr(col, 79), r.local_supplementary_allowance)
        _write(ws, addr(col, 80), r.house_rent_allowance)
        _write(ws, addr(col, 81), r.travel_allowance)
        _write(ws, addr(col, 82), r.other)

    tn_temp_records: List[models.PostStatus] = (
        db.query(models.PostStatus)
        .filter(models.PostStatus.district == "Thane", models.PostStatus.category == "Temporary")
        .all()
    )
    for r in tn_temp_records:
        col = filled_cols.get(r.class_type) if r.status == "Filled" else vacant_cols.get(r.class_type) if r.status == "Vacant" else None
        if not col:
            continue
        _write(ws, addr(col, 88), r.posts)
        _write(ws, addr(col, 89), r.salary)
        _write(ws, addr(col, 90), r.grade_pay)
        _write(ws, addr(col, 92), r.special_pay)
        _write(ws, addr(col, 93), r.dearness_allowance)
        _write(ws, addr(col, 94), r.local_supplementary_allowance)
        _write(ws, addr(col, 95), r.house_rent_allowance)
        _write(ws, addr(col, 96), r.travel_allowance)
        _write(ws, addr(col, 97), r.other)


def populate_post_expenses(wb, db: Session):
    sheet_name = ORIGINAL_SHEET_NAMES.get("post_expenses")
    if sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=500, detail=f"Sheet '{sheet_name}' not found in original workbook")
    ws = wb[sheet_name]
    def addr(col: str, row: int) -> str:
        return f"{col}{row}"

    def aggregate_counts(records: List[models.PostExpenses]) -> Dict[str, Dict[str, int]]:
        result: Dict[str, Dict[str, int]] = {}
        for r in records:
            cls = str(r.class_type)
            if cls not in result:
                result[cls] = {"filled": 0, "vacant": 0}
            result[cls]["filled"] += int(r.filled_posts or 0)
            result[cls]["vacant"] += int(r.vacant_posts or 0)
        return result

    def write_for_district(district: str, row_map: Dict[str, int], fixed_row: int):
        perm_records: List[models.PostExpenses] = (
            db.query(models.PostExpenses)
            .filter(models.PostExpenses.district == district, models.PostExpenses.category == "Permanent")
        .all()
    )
        temp_records: List[models.PostExpenses] = (
            db.query(models.PostExpenses)
            .filter(models.PostExpenses.district == district, models.PostExpenses.category == "Temporary")
        .all()
    )
        perm_counts = aggregate_counts(perm_records)
        temp_counts = aggregate_counts(temp_records)
        for cls, row in row_map.items():
            _write(ws, addr("C", row), perm_counts.get(cls, {}).get("filled"))
            _write(ws, addr("D", row), perm_counts.get(cls, {}).get("vacant"))
        for cls, row in row_map.items():
            _write(ws, addr("E", row), temp_counts.get(cls, {}).get("filled"))
            _write(ws, addr("F", row), temp_counts.get(cls, {}).get("vacant"))
        representative = (
            db.query(models.PostExpenses)
            .filter(models.PostExpenses.district == district)
            .first()
        )
        if representative:
            from config import POST_EXPENSES_DISTRICT_COMPONENT_FIELD
            active_component = POST_EXPENSES_DISTRICT_COMPONENT_FIELD.get(district)
            if active_component == 'SeventhPayCommissionDifferenceNPS':
                selected_value = representative.seventh_pay_commission_difference_nps
            elif active_component == 'NPS':
                selected_value = representative.nps
            elif active_component == 'SeventhPayCommissionDifference':
                selected_value = representative.seventh_pay_commission_difference
            else:
                selected_value = None
            fixed_values = [
                representative.medical_expenses,
                representative.festival_advance,
                representative.swagram_maharashtra_darshan,
                selected_value,
                representative.other,
            ]
            fixed_cols = ["C", "D", "E", "F", "G"]
            for col, val in zip(fixed_cols, fixed_values):
                _write(ws, addr(col, fixed_row), val)

    write_for_district("Mumbai City", {"1": 6, "2": 7, "3": 8, "4": 9}, 15)
    write_for_district("Mumbai Suburban", {"1": 22, "2": 23, "3": 24, "4": 25}, 31)
    write_for_district("Thane", {"1": 38, "2": 39, "3": 40, "4": 41}, 47)
    write_for_district("Palghar", {"1": 54, "2": 55, "3": 56, "4": 57}, 63)
    write_for_district("Raigad", {"1": 70, "2": 71, "3": 72, "4": 73}, 79)
    write_for_district("Ratnagiri", {"1": 86, "2": 87, "3": 88, "4": 89}, 95)
    write_for_district("Sindhudurg", {"1": 102, "2": 103, "3": 104, "4": 105}, 111)




def populate_unit_expenditure(wb, db: Session):
    sheet_name = ORIGINAL_SHEET_NAMES.get("unit_expenditure")
    if sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=500, detail=f"Sheet '{sheet_name}' not found in original workbook")
    ws = wb[sheet_name]
    def addr(col: str, row: int) -> str:
        return f"{col}{row}"

    from config import PRIMARY_UNITS
    col_order = ["C", "D", "E", "F", "G", "H", "I", "J", "K"]
    fields = [
        "expenditure_2021_22",
        "expenditure_2022_23",
        "expenditure_2023_24",
        "budget_2024_25",
        "forecast_2024_25",
        "budget_2025_26_estimating_officer",
        "budget_2025_26_controlling_officer",
        "budget_2025_26_admin_dept",
        "budget_2025_26_finance_dept",
    ]

    def write_district(district: str, start_row: int):
        row_map = {unit: row for unit, row in zip(PRIMARY_UNITS, range(start_row, start_row + len(PRIMARY_UNITS)))}
        items: List[models.UnitExpenditure] = (
            db.query(models.UnitExpenditure)
            .filter(models.UnitExpenditure.district == district)
        .all()
    )
        for it in items:
            row = row_map.get(it.unit_account)
            if not row:
                continue
            for col, field in zip(col_order, fields):
                _write(ws, addr(col, row), getattr(it, field, None))

    write_district("Mumbai City", 7)
    write_district("Mumbai Suburban", 30)
    write_district("Thane", 53)
    write_district("Palghar", 76)
    write_district("Raigad", 99)
    write_district("Ratnagiri", 122)
    write_district("Sindhudurg", 145)
